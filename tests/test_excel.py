# -*- coding: utf-8 -*-
"""Excel 导出测试：工作表结构、列格式、派生列与命令行入口。"""

from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from douban_top250.clean import split_values
from douban_top250.excel import (
    COLUMNS,
    as_float,
    as_int,
    build_workbook,
    columns_for,
    main,
    normalize_records,
    normalize_value,
    write_workbook,
)

openpyxl = pytest.importorskip("openpyxl")


class FakeNumpyInt:
    """模拟 numpy.int64：不是 int 但带 item()。"""

    def __init__(self, value):
        self._value = value

    def item(self):
        return self._value


def sheet_values(sheet):
    return [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]


def test_build_workbook_creates_two_sheets(sample_payload):
    workbook = build_workbook(sample_payload)

    assert workbook.sheetnames == ["Top250", "统计概览"]
    sheet = workbook["Top250"]
    assert sheet["A1"].value.startswith("豆瓣电影 Top250")
    assert "2026-01-01 00:00:00" in sheet["A1"].value


def test_data_sheet_layout_and_formats(sample_payload):
    sheet = build_workbook(sample_payload)["Top250"]
    headers = [sheet.cell(row=2, column=index).value for index in range(1, len(COLUMNS) + 1)]

    assert headers == [label for label, _, _, _, _ in COLUMNS]
    assert sheet.freeze_panes == "C3"
    assert sheet.auto_filter.ref == "A2:M5"          # 表头 1 行 + 3 条数据

    first = sheet.cell(row=3, column=1).value
    rating = sheet.cell(row=3, column=8).value
    votes = sheet.cell(row=3, column=9).value
    assert (first, rating, votes) == (1, 9.7, 3347760)
    assert sheet.cell(row=3, column=8).number_format == "0.0"
    assert sheet.cell(row=3, column=9).number_format == "#,##0"
    assert sheet.cell(row=3, column=12).hyperlink is not None    # 详情页链接
    assert sheet.column_dimensions["B"].width == 20
    assert sheet.cell(row=4, column=1).fill.fgColor.rgb.endswith("F2F5FA")   # 斑马纹


def test_extra_columns_only_added_when_present(sample_payload):
    base = [spec[1] for spec in columns_for(sample_payload["movies"])]
    assert base == [spec[1] for spec in COLUMNS]

    enriched = dict(sample_payload)
    enriched["movies"] = [dict(sample_payload["movies"][0], decade="1990 年代", votes_wan=334.8)]
    keys = [spec[1] for spec in columns_for(enriched["movies"])]

    assert "decade" in keys and "votes_wan" in keys
    assert keys[:len(COLUMNS)] == [spec[1] for spec in COLUMNS]


def test_stats_sheet_contains_sections_and_numbers(sample_payload):
    sheet = build_workbook(sample_payload)["统计概览"]
    values = sheet_values(sheet)

    for section in ("基本统计", "评分区间分布", "年代分布", "类型 Top10（可多类型重复计数）",
                    "国家/地区 Top10（合拍片重复计数）"):
        assert section in values

    assert "影片总数" in values and 3 in values
    assert round((9.7 + 9.6 + 8.1) / 3, 2) in values        # 平均评分
    # 类型按空白与斜杠拆分：剧情 / 犯罪 / 爱情 / 同性 / 喜剧
    assert "剧情" in values
    assert "中国香港" in values                             # 地区拆分不能把两个地区粘在一起


def test_multi_value_share_uses_all_occurrences(sample_payload):
    """类型/地区的占比分母应是"所有出现次数"，而不是前 10 名的合计。"""
    sheet = build_workbook(sample_payload)["统计概览"]
    rows = {row[0].value: (row[1].value, row[2].value) for row in sheet.iter_rows() if row[0].value}

    occurrences = [value for movie in sample_payload["movies"] for value in split_values(movie["genres"])]
    count, share = rows["剧情"]

    assert count == occurrences.count("剧情")
    assert share == round(count / len(occurrences), 4)


def test_write_workbook_produces_readable_file(tmp_path, sample_payload):
    path = write_workbook(sample_payload, tmp_path / "nested" / "out.xlsx")

    assert path.is_file() and path.stat().st_size > 0
    reloaded = load_workbook(str(path))
    assert reloaded["Top250"].max_row == 5
    assert reloaded["Top250"].cell(row=3, column=2).value == "肖申克的救赎"


def test_cli_main_generates_workbook(tmp_path, sample_payload):
    src = tmp_path / "in.json"
    src.write_text(json.dumps(sample_payload, ensure_ascii=False), encoding="utf-8")
    dst = tmp_path / "out.xlsx"

    assert main([str(src), str(dst)]) == 0
    assert dst.is_file()


def test_cli_main_reports_missing_input(tmp_path, capsys):
    assert main([str(tmp_path / "missing.json")]) == 1
    assert "输入文件不存在" in capsys.readouterr().err


def test_normalize_value_handles_numpy_like_and_nan():
    assert normalize_value(FakeNumpyInt(7)) == 7
    assert normalize_value(float("nan")) is None
    assert normalize_value(None) is None
    assert normalize_value(True) is True
    assert normalize_value("剧情") == "剧情"


def test_normalize_records_blanks_empty_strings():
    records = normalize_records([{"title": "   ", "quote": "好片", "votes": FakeNumpyInt(12)}])

    assert records[0]["title"] is None
    assert records[0]["quote"] == "好片"
    assert records[0]["votes"] == 12


def test_as_int_and_as_float_are_tolerant():
    assert as_int("1,234") == 1234
    assert as_int(None) == 0
    assert as_int("abc") == 0
    assert as_float("9.7") == pytest.approx(9.7)
    assert as_float("") == 0.0
