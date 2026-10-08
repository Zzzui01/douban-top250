# -*- coding: utf-8 -*-
"""Excel 导出层：把抓取（或清洗后）的记录写成带格式的工作簿。

工作表：
  1. Top250   —— 完整数据（标题行、冻结表头、自动筛选、超链接、数字格式、斑马纹）
  2. 统计概览 —— 基本统计 / 评分区间分布 / 年代分布 / 类型 Top10 / 国家地区 Top10

若记录里含清洗阶段派生的列（decade、rating_band、votes_wan 等），会自动追加到数据表，
因此既可以直接导出原始结果，也可以导出 Pandas 清洗后的结果。
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple, Union

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .clean import RATING_BANDS, split_values
from .config import JSON_NAME, OUTPUT_DIR, XLSX_NAME
from .logging_setup import get_logger
from .models import to_native

logger = get_logger(__name__)

PathLike = Union[str, Path]

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
BAND_FILL = PatternFill("solid", fgColor="F2F5FA")
SECTION_FONT = Font(bold=True, color="FFFFFF", size=11)
SECTION_FILL = PatternFill("solid", fgColor="548235")
SUB_FONT = Font(bold=True, color="1F3864")
LINK_FONT = Font(color="0563C1", underline="single", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

#: (中文表头, 字段名, 列宽, 对齐, 数字格式)
COLUMNS: Sequence[Tuple[str, str, int, str, Optional[str]]] = (
    ("排名", "rank", 6, "center", None),
    ("片名", "title", 20, "left", None),
    ("外文名", "foreign_title", 30, "left", None),
    ("别名", "aliases", 30, "left", None),
    ("年份", "year", 8, "center", None),
    ("国家/地区", "country", 18, "left", None),
    ("类型", "genres", 20, "left", None),
    ("评分", "rating", 8, "center", "0.0"),
    ("评价人数", "votes", 12, "right", "#,##0"),
    ("短评", "quote", 40, "left", None),
    ("导演/主演", "staff", 46, "left", None),
    ("详情页链接", "url", 42, "left", None),
    ("海报链接", "cover", 42, "left", None),
)

#: 清洗阶段派生列，仅在记录里存在时追加
EXTRA_COLUMNS: Sequence[Tuple[str, str, int, str, Optional[str]]] = (
    ("年代", "decade", 12, "center", None),
    ("评分档位", "rating_band", 14, "center", None),
    ("类型数", "genre_count", 8, "center", "0"),
    ("地区数", "country_count", 8, "center", "0"),
    ("评价人数(万)", "votes_wan", 13, "right", "0.0"),
    ("有短评", "has_quote", 9, "center", None),
)

WRAP_KEYS = {"aliases", "quote", "staff", "foreign_title"}
LINK_KEYS = ("url", "cover")


# --------------------------------------------------------------------------- #
# 取值与格式化
# --------------------------------------------------------------------------- #
def normalize_value(value: object) -> object:
    """把 numpy/pandas 标量转成 openpyxl 能写的原生类型（NaN 转 None）。"""
    return to_native(value)


def normalize_records(records: Sequence[Dict[str, object]]) -> List[Dict[str, object]]:
    """逐字段归一化记录，顺带把空字符串转成 None（Excel 里显示为空单元格）。"""
    cleaned = []
    for record in records:
        row = {}
        for key, value in dict(record).items():
            value = normalize_value(value)
            if isinstance(value, str) and not value.strip():
                value = None
            row[key] = value
        cleaned.append(row)
    return cleaned


def columns_for(records: Sequence[Dict[str, object]]) -> List[Tuple[str, str, int, str, Optional[str]]]:
    """基础列 + 记录中实际存在的派生列。"""
    columns = list(COLUMNS)
    keys = set()
    for record in records:
        keys.update(record.keys())
    for spec in EXTRA_COLUMNS:
        if spec[1] in keys:
            columns.append(spec)
    return columns


def as_int(value: object, default: int = 0) -> int:
    try:
        return int(float(str(value).replace(",", "")))
    except (TypeError, ValueError):
        return default


def as_float(value: object, default: float = 0.0) -> float:
    try:
        return float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return default


def style_header(sheet, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = sheet.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 22


def style_section(sheet, row: int, span: int) -> None:
    sheet.cell(row=row, column=1).font = SECTION_FONT
    sheet.cell(row=row, column=1).fill = SECTION_FILL
    for col in range(2, span + 1):
        sheet.cell(row=row, column=col).fill = SECTION_FILL
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)


# --------------------------------------------------------------------------- #
# 工作表构建
# --------------------------------------------------------------------------- #
def build_data_sheet(workbook: Workbook, payload: Dict[str, object]):
    records = normalize_records(payload.get("movies") or [])
    columns = columns_for(records)
    ncols = len(columns)

    sheet = workbook.active
    sheet.title = "Top250"
    title = "豆瓣电影 Top250（来源：{}，抓取时间：{}，共 {} 条）".format(
        payload.get("source") or "未知", payload.get("crawled_at") or "未知", len(records)
    )
    sheet.cell(row=1, column=1, value=title).font = TITLE_FONT
    sheet.cell(row=1, column=1).alignment = Alignment(vertical="center")
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 1))
    sheet.row_dimensions[1].height = 26

    header = 2
    for idx, (label, _, _, _, _) in enumerate(columns, start=1):
        sheet.cell(row=header, column=idx, value=label)
    style_header(sheet, header, ncols)

    for offset, record in enumerate(records):
        row = header + 1 + offset
        for idx, (_, key, _, align, numfmt) in enumerate(columns, start=1):
            value = record.get(key)
            if key in ("votes", "rank", "year", "genre_count", "country_count"):
                value = as_int(value) if value is not None else None
            elif key in ("rating", "votes_wan"):
                value = as_float(value) if value is not None else None
            cell = sheet.cell(row=row, column=idx, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=key in WRAP_KEYS)
            if numfmt:
                cell.number_format = numfmt
            if key in LINK_KEYS and value:
                cell.hyperlink = value
                cell.font = LINK_FONT
            if offset % 2 == 1:
                cell.fill = BAND_FILL

    for idx, (_, _, width, _, _) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = sheet.cell(row=header + 1, column=3)
    sheet.auto_filter.ref = "A{0}:{1}{2}".format(header, get_column_letter(ncols), header + len(records))
    return sheet


def _decade(year: int) -> str:
    if year <= 0:
        return "未知年代"
    if year < 1990:
        return "1990 年以前"
    return "{} 年代".format(year // 10 * 10)


def _band(rating: float) -> str:
    if rating >= 9.5:
        return RATING_BANDS[0]
    if rating >= 9.0:
        return RATING_BANDS[1]
    if rating >= 8.5:
        return RATING_BANDS[2]
    return RATING_BANDS[3]


def build_stats_sheet(workbook: Workbook, payload: Dict[str, object]):
    records = normalize_records(payload.get("movies") or [])
    sheet = workbook.create_sheet("统计概览")
    for column, width in (("A", 24), ("B", 16), ("C", 14), ("D", 14), ("E", 14)):
        sheet.column_dimensions[column].width = width

    ratings = [as_float(m.get("rating")) for m in records if m.get("rating") is not None]
    votes = [as_int(m.get("votes")) for m in records if m.get("votes") is not None]
    years = [as_int(m.get("year")) for m in records if as_int(m.get("year")) > 0]

    row = 1
    sheet.cell(row=row, column=1, value="豆瓣 Top250 统计概览").font = TITLE_FONT
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    sheet.row_dimensions[row].height = 24
    row += 2

    sheet.cell(row=row, column=1, value="基本统计")
    style_section(sheet, row, 2)
    row += 1

    basics = [
        ("影片总数", len(records), "0"),
        ("平均评分", round(sum(ratings) / len(ratings), 2) if ratings else 0, "0.00"),
        ("最高评分", max(ratings) if ratings else 0, "0.0"),
        ("最低评分", min(ratings) if ratings else 0, "0.0"),
        ("评分 ≥ 9.0 部数", sum(1 for r in ratings if r >= 9.0), "0"),
        ("评价人数合计", sum(votes), "#,##0"),
        ("平均评价人数", round(sum(votes) / len(votes)) if votes else 0, "#,##0"),
        ("最早年份", min(years) if years else 0, "0"),
        ("最新年份", max(years) if years else 0, "0"),
        ("数据抓取时间", payload.get("crawled_at") or "未知", None),
    ]
    if payload.get("cleaned_at"):
        basics.append(("清洗时间", payload.get("cleaned_at"), None))
    report = payload.get("report")
    if isinstance(report, dict):
        basics.append(("清洗后记录数", report.get("total_out", len(records)), "0"))
        basics.append(("清洗剔除记录数", report.get("dropped", 0), "0"))

    for name, value, fmt in basics:
        sheet.cell(row=row, column=1, value=name).font = SUB_FONT
        cell = sheet.cell(row=row, column=2, value=value)
        cell.alignment = Alignment(horizontal="left")
        if fmt:
            cell.number_format = fmt
        row += 1
    row += 1

    genre_counter = Counter(
        genre for m in records for genre in split_values(m.get("genres"))
    )
    country_counter = Counter(
        country for m in records for country in split_values(m.get("country"))
    )

    decade_counter = Counter(_decade(year) for year in years)
    decade_order = ["1990 年以前"] + ["{} 年代".format(y) for y in range(1990, 2030, 10)]

    blocks = (
        ("评分区间分布", Counter(_band(r) for r in ratings),
         [b for b in RATING_BANDS if any(_band(r) == b for r in ratings)], len(ratings)),
        ("年代分布", decade_counter,
         [d for d in decade_order if decade_counter.get(d)], len(years)),
        # 类型/地区是多值字段：占比的分母取"所有出现次数"，而不是前 10 名的合计
        ("类型 Top10（可多类型重复计数）", genre_counter, None, sum(genre_counter.values())),
        ("国家/地区 Top10（合拍片重复计数）", country_counter, None, sum(country_counter.values())),
    )

    for block_title, counter, fixed_order, total in blocks:
        sheet.cell(row=row, column=1, value=block_title)
        style_section(sheet, row, 3)
        row += 1

        for col, name in enumerate(("名称", "数量", "占比"), start=1):
            cell = sheet.cell(row=row, column=col, value=name)
            cell.font = SUB_FONT
            cell.border = BORDER
        row += 1

        items = counter.most_common()
        if fixed_order is None:
            items = items[:10]
        else:
            order = {name: index for index, name in enumerate(fixed_order)}
            items = sorted(items, key=lambda kv: order.get(kv[0], 999))

        denominator = total or sum(count for _, count in items) or 1
        for name, count in items:
            sheet.cell(row=row, column=1, value=name).border = BORDER
            count_cell = sheet.cell(row=row, column=2, value=count)
            count_cell.number_format = "0"
            count_cell.border = BORDER
            share_cell = sheet.cell(row=row, column=3, value=round(count / denominator, 4))
            share_cell.number_format = "0.0%"
            share_cell.border = BORDER
            row += 1

        if fixed_order is not None and items and total:
            sheet.cell(row=row, column=1, value="合计").font = SUB_FONT
            total_cell = sheet.cell(row=row, column=2, value=sum(count for _, count in items))
            total_cell.font = SUB_FONT
            share_cell = sheet.cell(row=row, column=3, value=sum(count for _, count in items) / denominator)
            share_cell.number_format = "0.0%"
            share_cell.font = SUB_FONT
            row += 1
        row += 1

    return sheet


def build_workbook(payload: Dict[str, object]) -> Workbook:
    """把结果 payload 转成 Workbook 对象（不落盘，便于测试）。"""
    workbook = Workbook()
    build_data_sheet(workbook, payload)
    build_stats_sheet(workbook, payload)
    logger.debug("工作簿构建完成：%d 条记录", len(payload.get("movies") or []))
    return workbook


def write_workbook(payload: Dict[str, object], path: PathLike = OUTPUT_DIR / XLSX_NAME) -> Path:
    """生成并保存工作簿，返回输出路径。"""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook(payload)
    workbook.save(str(target))
    logger.info("已保存 Excel -> %s（%d 条）", target, len(payload.get("movies") or []))
    return target


# --------------------------------------------------------------------------- #
# 命令行入口：python -m douban_top250.excel [输入.json] [输出.xlsx]
# --------------------------------------------------------------------------- #
def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="把 JSON 结果转换为带格式的 Excel 工作簿")
    parser.add_argument("src", nargs="?", default=str(OUTPUT_DIR / JSON_NAME), help="输入 JSON 路径")
    parser.add_argument("dst", nargs="?", default=None, help="输出 XLSX 路径，默认与输入同目录")
    args = parser.parse_args(argv)

    src = Path(args.src)
    dst = Path(args.dst) if args.dst else src.parent / XLSX_NAME
    if not src.is_file():
        print("输入文件不存在：{}".format(src), file=sys.stderr)
        return 1
    try:
        with open(str(src), encoding="utf-8") as fp:
            payload = json.load(fp)
        if isinstance(payload, list):
            payload = {"source": "未知", "crawled_at": None, "movies": payload}
        path = write_workbook(payload, dst)
    except (OSError, ValueError) as exc:
        print("生成失败：{}: {}".format(type(exc).__name__, exc), file=sys.stderr)
        return 1
    print("已生成: {}（{} 字节，共 {} 条）".format(path, path.stat().st_size, len(payload.get("movies") or [])))
    return 0


if __name__ == "__main__":  # pragma: no cover
    sys.exit(main())


__all__ = [
    "COLUMNS",
    "EXTRA_COLUMNS",
    "RATING_BANDS",
    "build_workbook",
    "write_workbook",
    "build_data_sheet",
    "build_stats_sheet",
    "columns_for",
    "normalize_records",
    "normalize_value",
    "as_int",
    "as_float",
    "main",
]
